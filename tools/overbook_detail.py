# -*- coding: utf-8 -*-
"""打印若干「原文时长 > 时段长度」格子的明细行，看清数字是怎么来的"""
import io, os, re, sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

XLSX = r"D:\UsrFiles\Documents\NCUK IFY Self Study\02_学习与教材\Schedule.xlsx"
OUT = os.path.join(os.environ["TEMP"], "overbook_detail.txt")

TARGETS = {("3", "17"), ("3", "31"), ("1", "20"), ("12", "16"), ("11", "17"),
           ("3", "18"), ("12", "30"), ("4", "15"), ("3", "6"), ("12", "25")}

DAY_RE = re.compile(r"^(\d{1,2})月(\d{1,2})日")
TIME_RE = re.compile(r"^\s*(\d{1,2}):(\d{2})\s*[-–—]\s*(\d{1,2}):(\d{2})\s*(.*?)\s*$")
DASH_RE = re.compile(r"^[·•]\s*(.*)$")
TOTAL_RE = re.compile(r"共\s*(\d+)\s*分钟")
PER_RE = re.compile(r"(?:每节|各|每个|每小块)\s*(\d+)\s*分钟")
MIN_RE = re.compile(r"(\d+)\s*[′'’]")
MIN2_RE = re.compile(r"(\d+)\s*分钟")
SEC_RE = re.compile(r"§\s*(\d+\.\d+)")


def label_of(t):
    i = t.find("：")
    return t[:i] if i > 0 else t


def stated(t):
    seg = label_of(t)
    m = TOTAL_RE.search(seg)
    if m: return int(m.group(1))
    p = PER_RE.search(seg)
    if p:
        n = int(p.group(1)); c = len(set(SEC_RE.findall(seg)))
        if c >= 1: return n * c
    m = MIN_RE.search(seg) or MIN2_RE.search(seg)
    return int(m.group(1)) if m else 0


out = []
wb = openpyxl.load_workbook(XLSX, data_only=True)
for name in wb.sheetnames:
    if not (len(name) == 6 and name.isdigit()):
        continue
    ws = wb[name]
    for row in ws.iter_rows(values_only=True):
        for c in row:
            if not isinstance(c, str) or "月" not in c:
                continue
            lines = [l.replace("\r", "") for l in c.split("\n")]
            m0 = DAY_RE.match(lines[0].strip())
            if not m0:
                continue
            if (m0.group(1), m0.group(2)) not in TARGETS:
                continue
            cur = None
            for ln in lines[1:]:
                t = ln.strip()
                if not t or t.strip("─-—=_").strip() == "":
                    continue
                m = TIME_RE.match(ln)
                if m:
                    cur = {"h": m.group(0).strip(), "body": []}
                    out.append(cur)
                elif cur is not None:
                    cur["body"].append(t)
            out.append({"h": "=========== " + lines[0].strip(), "body": []})

for blk in out:
    if blk["h"].startswith("====="):
        out2 = blk
        print(blk["h"], file=io.open(OUT, "a", encoding="utf-8"))
        continue
    det = []
    for ln in blk["body"][1:]:
        d = DASH_RE.match(ln)
        if d and d.group(1).strip():
            det.append(d.group(1).strip())
    fixed = sum(stated(d) for d in det)
    with io.open(OUT, "a", encoding="utf-8") as f:
        f.write("\n" + blk["h"] + "\n")
        if blk["body"]:
            f.write("  标题: " + blk["body"][0] + "\n")
        for d in det:
            f.write("   %3d'  %s\n" % (stated(d), d))
        f.write("   小计 %d'\n" % fixed)

print("->", OUT)
