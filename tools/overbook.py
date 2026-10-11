# -*- coding: utf-8 -*-
"""
量一遍 Schedule.xlsx 里「课表原文写死的分钟数加起来 > 时段长度」的格子，
算出每个格子需要延长多少分钟、后面还留了多少空隙。

解析规则必须和 DailyPlan.StatedMinutes 一致（只认第一个中文冒号之前的「标签段」）。
"""
import io, os, re, sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

XLSX = r"D:\UsrFiles\Documents\NCUK IFY Self Study\02_学习与教材\Schedule.xlsx"
OUT = os.path.join(os.environ["TEMP"], "overbook.txt")

DAY_RE = re.compile(r"^(\d{1,2})月(\d{1,2})日")
TIME_RE = re.compile(r"^\s*(\d{1,2}):(\d{2})\s*[-–—]\s*(\d{1,2}):(\d{2})\s*(.*?)\s*$")
DASH_RE = re.compile(r"^[·•]\s*(.*)$")
NEW_RE = re.compile(r"^新课\s+([A-Za-z0-9/]+)\s+(Ch\d+)")
PRAC_RE = re.compile(r"^练习\s+([A-Za-z0-9/]+)\s+(Ch\d+)")

TOTAL_RE = re.compile(r"共\s*(\d+)\s*分钟")
PER_RE = re.compile(r"(?:每节|各|每个|每小块)\s*(\d+)\s*分钟")
MIN_RE = re.compile(r"(\d+)\s*[′'’]")
MIN2_RE = re.compile(r"(\d+)\s*分钟")
SEC_RE = re.compile(r"§\s*(\d+\.\d+)")


def label_of(text):
    i = text.find("：")
    return text[:i] if i > 0 else text


def stated_minutes(text):
    seg = label_of(text)
    m = TOTAL_RE.search(seg)
    if m:
        return int(m.group(1))
    p = PER_RE.search(seg)
    if p:
        n = int(p.group(1))
        cnt = len(set(SEC_RE.findall(seg)))
        if cnt >= 1:
            return n * cnt
    m = MIN_RE.search(seg)
    if not m:
        m = MIN2_RE.search(seg)
    return int(m.group(1)) if m else 0


def tomin(h, m):
    return int(h) * 60 + int(m)


def hhmm(total):
    return "%02d:%02d" % (total // 60, total % 60)


rows = []
wb = openpyxl.load_workbook(XLSX, data_only=True)
for name in wb.sheetnames:
    if not (len(name) == 6 and name.isdigit()):
        continue
    ws = wb[name]
    for row in ws.iter_rows(values_only=True):
        for c in row:
            if not isinstance(c, str) or "月" not in c:
                continue
            lines = [l.replace("\r", "") for l in c.split("\n")]
            if not DAY_RE.match(lines[0].strip()):
                continue
            slots = []
            cur = None
            for ln in lines[1:]:
                t = ln.strip()
                if not t or t.strip("─-—=_").strip() == "":
                    continue
                m = TIME_RE.match(ln)
                if m:
                    cur = {"date": lines[0].strip(), "sheet": name,
                           "start": tomin(m.group(1), m.group(2)),
                           "end": tomin(m.group(3), m.group(4)),
                           "subject": (m.group(5) or "").strip(), "body": []}
                    slots.append(cur)
                elif cur is not None:
                    cur["body"].append(t)
            for i, s in enumerate(slots):
                nxt = slots[i + 1]["start"] if i + 1 < len(slots) else None
                s["next"] = nxt
            rows.extend(slots)

over = []
for s in rows:
    if not s["body"]:
        continue
    title = s["body"][0]
    if NEW_RE.match(title) or PRAC_RE.match(title):
        continue                      # 这两类走别的分支，时长来自 exercises.tsv
    det = []
    for ln in s["body"][1:]:
        m = DASH_RE.match(ln)
        if m and m.group(1).strip():
            det.append(m.group(1).strip())
    if not det:
        continue
    fixed = sum(stated_minutes(d) for d in det)
    dur = s["end"] - s["start"]
    if fixed > dur:
        gap = (s["next"] - s["end"]) if s["next"] is not None else None
        over.append({
            "date": s["date"], "sheet": s["sheet"],
            "start": s["start"], "end": s["end"], "dur": dur,
            "subject": s["subject"], "title": title,
            "fixed": fixed, "need": fixed - dur, "gap": gap,
            "next": s["next"],
        })

out = []
out.append("原文时长 > 时段长度 的格子：%d 个" % len(over))
out.append("")
out.append("日期            时段            科目        需要 现长 差  后面空隙  原文和")
tot_need = 0
tight = 0
for o in over:
    tot_need += o["need"]
    g = "—" if o["gap"] is None else "%d′" % o["gap"]
    if o["gap"] is not None and o["gap"] < o["need"]:
        tight += 1
        g += " ⚠紧"
    out.append("%-14s %s-%s  %-10s %3d′ %3d′ +%d′  %-8s %3d′" % (
        o["date"], hhmm(o["start"]), hhmm(o["end"]), o["subject"][:10],
        o["fixed"], o["dur"], o["need"], g, o["fixed"]))
out.append("")
out.append("合计需要延长 %d 分钟（约 %.1f 小时）；其中后面空隙不够、会撞到下一格的：%d 个"
           % (tot_need, tot_need / 60.0, tight))
gaps = [o["gap"] for o in over if o["gap"] is not None]
if gaps:
    out.append("可用的下一格间隔：最小 %d′，最大 %d′" % (min(gaps), max(gaps)))
need_dist = {}
for o in over:
    need_dist[o["need"]] = need_dist.get(o["need"], 0) + 1
out.append("需要延长的分钟数分布：" + "、".join(
    "%d′×%d" % (k, v) for k, v in sorted(need_dist.items())))

io.open(OUT, "w", encoding="utf-8", newline="").write("\n".join(out))
print("->", OUT)
